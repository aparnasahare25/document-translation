import os
import sys
import time
from typing import Optional

# Add project root to sys.path so 'scripts' package is importable
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from scripts.translator_service_pptx import TranslatorService

# Target-language display font: PowerPoint keeps separate font slots for
# Latin vs. East Asian (and complex-script) text, so translated CJK text
# won't actually render in the right typeface unless <a:ea> is set
# explicitly alongside <a:latin> (see word_translation_service.py for the
# same rationale on the Word/OOXML side).
_TARGET_FONT_BY_LANG_PREFIX = {
    "en": "Arial",
    "ja": "MS Gothic",
}


def _resolve_target_font(to_lang: str) -> Optional[str]:
    """Map a target language code to a display font.

    Returns None (leave the presentation's original font untouched) for any
    target language other than English/Japanese.
    """
    t = (to_lang or "").strip().lower()
    for prefix, font_name in _TARGET_FONT_BY_LANG_PREFIX.items():
        if t.startswith(prefix):
            return font_name
    return None


class PptxFormatter:
    """Translate and format a PPTX file using batch Azure Translator calls.

    Mirrors services/word_translation_service.py's DocxFormatter: the same
    collect -> translate (batched 3-stage pipeline) -> redistribute -> format
    -> save flow, adapted for PowerPoint's shape/slide model instead of
    Word's paragraph/section model.
    """

    def __init__(self, from_lang: str, to_lang: str):
        self.from_lang = from_lang
        self.to_lang = to_lang
        self.translator = TranslatorService()

    # ------------------------------------------------------------------
    # Helper: collect (runs, full_text) groups from a single text frame
    # ------------------------------------------------------------------
    @staticmethod
    def _collect_text_frame_groups(text_frame, groups: list):
        for paragraph in text_frame.paragraphs:
            runs = list(paragraph.runs)
            if runs:
                full_text = "".join(run.text for run in runs)
                groups.append((runs, full_text))

    # ------------------------------------------------------------------
    # Helper: apply post-translation formatting to a single text frame
    # ------------------------------------------------------------------
    @staticmethod
    def _format_text_frame_structure(text_frame, target_font: Optional[str] = None):
        """
        Apply structural/visual formatting to a text frame's runs AFTER text
        has already been translated and written back.

        - Resets character spacing to normal (avoids layout issues with
          CJK / wide-character scripts), mirroring DocxFormatter's approach.
        - Forces the run's font family to `target_font` (Arial for English,
          MS Gothic for Japanese) across the Latin, East Asian, and
          complex-script slots, so translated text renders correctly
          regardless of what font the source deck originally used.
        - Font-size is left untouched (original sizes preserved), consistent
          with the Word workflow's choice not to force a fixed size.
        """
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                try:
                    if run.font is not None and run.font._element is not None:
                        run.font._element.set("spc", "0")
                except Exception as e:
                    print(f"[format] Error adjusting character spacing: {e}")

                if target_font:
                    try:
                        run.font.name = target_font  # sets <a:latin typeface=...>
                        rPr = run.font._element
                        latin = rPr.find(qn("a:latin"))
                        # <a:ea> (East Asian) and <a:cs> (complex script) must
                        # directly follow <a:latin> in the OOXML sequence.
                        insert_after = latin
                        for tag in ("a:ea", "a:cs"):
                            elem = rPr.find(qn(tag))
                            if elem is None:
                                elem = rPr.makeelement(qn(tag), {})
                                insert_after.addnext(elem)
                            elem.set("typeface", target_font)
                            insert_after = elem
                    except Exception as e:
                        print(f"[format] Error setting font family: {e}")

    # ------------------------------------------------------------------
    # Helper: walk a shape tree (recursing into grouped shapes), collecting
    # translatable (runs, full_text) groups from text boxes, tables, and
    # chart titles/axis titles.
    # ------------------------------------------------------------------
    def _walk_shapes_collect(self, shapes, groups: list):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                self._walk_shapes_collect(shape.shapes, groups)
                continue

            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    for cell in row.cells:
                        self._collect_text_frame_groups(cell.text_frame, groups)
                continue

            if shape.has_chart:
                chart = shape.chart
                try:
                    if chart.has_title and chart.chart_title.has_text_frame:
                        self._collect_text_frame_groups(chart.chart_title.text_frame, groups)
                except Exception as e:
                    print(f"[collect] Error reading chart title: {e}")
                for axis_attr in ("category_axis", "value_axis"):
                    try:
                        axis = getattr(chart, axis_attr, None)
                        if axis is not None and axis.has_title and axis.axis_title.has_text_frame:
                            self._collect_text_frame_groups(axis.axis_title.text_frame, groups)
                    except Exception as e:
                        print(f"[collect] Error reading {axis_attr} title: {e}")
                continue

            if shape.has_text_frame:
                self._collect_text_frame_groups(shape.text_frame, groups)

    # ------------------------------------------------------------------
    # Helper: same walk as above, but applies formatting instead of collecting
    # ------------------------------------------------------------------
    def _walk_shapes_format(self, shapes, target_font: Optional[str] = None):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                self._walk_shapes_format(shape.shapes, target_font)
                continue

            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    for cell in row.cells:
                        self._format_text_frame_structure(cell.text_frame, target_font)
                        # Allow the (likely wider, translated) text to wrap
                        # instead of overflowing the cell.
                        try:
                            cell.text_frame.word_wrap = True
                        except Exception as e:
                            print(f"[format] Error enabling table cell word wrap: {e}")
                continue

            if shape.has_chart:
                chart = shape.chart
                try:
                    if chart.has_title and chart.chart_title.has_text_frame:
                        self._format_text_frame_structure(chart.chart_title.text_frame, target_font)
                except Exception:
                    pass
                for axis_attr in ("category_axis", "value_axis"):
                    try:
                        axis = getattr(chart, axis_attr, None)
                        if axis is not None and axis.has_title and axis.axis_title.has_text_frame:
                            self._format_text_frame_structure(axis.axis_title.text_frame, target_font)
                    except Exception:
                        pass
                continue

            if shape.has_text_frame:
                self._format_text_frame_structure(shape.text_frame, target_font)

    # ------------------------------------------------------------------
    # Pass 1 – Collect: gather ALL translatable paragraph groups across the
    # presentation (slides, tables, charts, grouped shapes, speaker notes)
    # ------------------------------------------------------------------
    def _collect_paragraph_groups(self, prs):
        """
        Returns a list of paragraph groups.
        Each group = (list_of_runs, full_text)
        """
        groups: list = []

        for slide in prs.slides:
            self._walk_shapes_collect(slide.shapes, groups)

            # Speaker notes (analogous to Word's headers/footers: separate
            # content area that's easy to miss if not handled explicitly).
            if slide.has_notes_slide:
                self._collect_text_frame_groups(slide.notes_slide.notes_text_frame, groups)

        return groups

    # ------------------------------------------------------------------
    # Pass 3 – Format: apply visual formatting to the whole presentation
    # ------------------------------------------------------------------
    def _apply_formatting(self, prs):
        target_font = _resolve_target_font(self.to_lang)
        for slide in prs.slides:
            self._walk_shapes_format(slide.shapes, target_font)
            if slide.has_notes_slide:
                self._format_text_frame_structure(slide.notes_slide.notes_text_frame, target_font)

    def _redistribute_text_to_runs(self, runs, translated_text):
        """
        Distribute translated text back to runs based on
        ORIGINAL CHARACTER PROPORTIONS (not tokens).

        Identical approach to DocxFormatter — python-pptx runs expose the
        same plain `run.text` get/set property as python-docx runs, so the
        proportional-allocation algorithm carries over unchanged.
        """

        total_src_len = sum(len(run.text) for run in runs)

        if total_src_len == 0:
            return

        total_trans_len = len(translated_text)
        cursor = 0

        for i, run in enumerate(runs):
            src_len = len(run.text)

            if i == len(runs) - 1:
                # Last run gets everything remaining
                run.text = translated_text[cursor:]
            else:
                # Proportional allocation
                proportion = src_len / total_src_len
                alloc_len = int(round(proportion * total_trans_len))

                run.text = translated_text[cursor: cursor + alloc_len]
                cursor += alloc_len

    # ------------------------------------------------------------------
    # Main entry point
    # ------------------------------------------------------------------
    def process_pptx(self, input_pptx_path: str, output_pptx_path: str) -> float:
        """
        Translate and format a PPTX file.

        Strategy
        --------
        1. Load the presentation.
        2. COLLECT  – walk all slides (text boxes, tables, charts, grouped
                      shapes, speaker notes) and gather (run_refs, text) pairs
                      into flat lists.
        3. TRANSLATE – send all texts through the full 3-stage pipeline
                       (MT -> LLM1 -> RAG+LLM2), same as the Word workflow.
        4. WRITE BACK – assign each translated string back to its runs via the
                        stored references, proportionally by character length.
        5. FORMAT   – apply structural fixes (character spacing, table cell
                      word-wrap) now that the correct text is in place.
        6. Save.
        """
        start_time = time.time()
        print("Started PPTX formatting.")

        # Step 1 – Load
        prs = Presentation(input_pptx_path)

        # Step 2 – Collect
        print("Collecting paragraph groups...")
        groups = self._collect_paragraph_groups(prs)
        print(f"  → {len(groups)} paragraphs found.")

        # Extract full paragraph texts
        paragraph_texts = [full_text for (_, full_text) in groups]

        # Step 3 – Batch translate: full 3-stage pipeline (MT → LLM1 → RAG+LLM2)
        print(f"Translating paragraphs '{self.from_lang}' → '{self.to_lang}' ...")
        translated_paragraphs = self.translator.batch_translate_with_pipeline(
            paragraph_texts,
            self.from_lang,
            self.to_lang
        )

        # Step 4 – Write back
        print("Redistributing translated text back to runs...")
        for (runs, _), translated in zip(groups, translated_paragraphs):
            if translated:
                self._redistribute_text_to_runs(runs, translated)

        # Step 5 – Apply formatting
        print("Applying presentation formatting...")
        self._apply_formatting(prs)

        # Step 6 – Save
        prs.save(output_pptx_path)

        end_time = time.time()
        duration = end_time - start_time
        print(f"\nDone! Presentation saved to: {output_pptx_path}")
        print(f"Total time taken: {duration:.2f} seconds")
        return duration


# ---------------------------------------------------------------------------
# Example usage
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    from_lang = "ja"
    to_lang = "en"
    input_pptx_path = r"C:\Users\Admin\Desktop\PPT Translation\Japanese to English Translation\JPN - SC委員会資料 - 10 pages.pptx"
    file_extension = os.path.splitext(input_pptx_path)[1]
    input_pptx_name = os.path.splitext(os.path.basename(input_pptx_path))[0]
    output_pptx_path = f"{input_pptx_name}_{from_lang}_{to_lang}{file_extension}"

    formatter = PptxFormatter(from_lang, to_lang)
    formatter.process_pptx(input_pptx_path, output_pptx_path)
